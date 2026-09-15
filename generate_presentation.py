from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).parent
PICTURES = ROOT / "backend" / "src" / "main" / "resources" / "pictures"
OUTPUT = ROOT / "Oracle-Migration-Studio-Overview.pptx"

W, H = 13.333, 7.5
NAVY = RGBColor(25, 17, 28)
BURGUNDY = RGBColor(50, 12, 27)
RED = RGBColor(236, 20, 58)
CORAL = RGBColor(255, 100, 91)
CREAM = RGBColor(248, 246, 242)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(26, 33, 38)
MUTED = RGBColor(102, 112, 119)
LINE = RGBColor(218, 215, 211)
TEAL = RGBColor(35, 139, 137)
GOLD = RGBColor(236, 177, 77)

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]


def box(slide, x, y, w, h, fill=None, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill or WHITE
    shape.line.color.rgb = line or (fill or WHITE)
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False, font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return tb


def title(slide, kicker, heading, sub=None, dark=False):
    text(slide, kicker.upper(), 0.72, 0.48, 4.5, 0.25, 10, RED if not dark else CORAL, True, "Aptos", italic=False)
    text(slide, heading, 0.70, 0.82, 11.6, 0.65, 27, WHITE if dark else INK, True, "Georgia")
    if sub:
        text(slide, sub, 0.72, 1.55, 10.8, 0.4, 13, RGBColor(216, 216, 212) if dark else MUTED)


def footer(slide, number, dark=False):
    text(slide, "TECH MAHINDRA  /  MIGRATION STUDIO", 0.72, 7.13, 5.5, 0.18, 8, RGBColor(183, 180, 179) if dark else MUTED, True)
    text(slide, f"{number:02d}", 12.0, 7.08, 0.6, 0.2, 9, CORAL if dark else RED, True, "Aptos", PP_ALIGN.RIGHT)


def add_image(slide, filename, x, y, w, h):
    slide.shapes.add_picture(str(PICTURES / filename), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def pill(slide, label, x, y, w, fill=RED, color=WHITE):
    box(slide, x, y, w, 0.32, fill, fill, True)
    text(slide, label, x, y + 0.055, w, 0.16, 8.5, color, True, "Aptos", PP_ALIGN.CENTER)


def bullet(slide, value, x, y, w, color=INK, accent=RED, size=14):
    box(slide, x, y + 0.08, 0.11, 0.11, accent, accent, True)
    text(slide, value, x + 0.22, y, w - 0.22, 0.38, size, color)


# 1. Cover
s = prs.slides.add_slide(blank)
box(s, 0, 0, W, H, NAVY)
box(s, 0, 0, 0.2, H, RED)
box(s, 8.65, 0, 4.68, H, BURGUNDY)
text(s, "ORACLE\nMIGRATION", 0.75, 0.78, 5.7, 1.28, 35, WHITE, True, "Georgia")
text(s, "MIGRATION STUDIO", 0.78, 2.35, 4.5, 0.3, 12, CORAL, True)
text(s, "A controlled path from SAP Sybase ASE metadata and data\nto ordered Oracle migration scripts and evidence-based reconciliation.", 0.78, 3.1, 6.25, 0.8, 18, RGBColor(220, 218, 213))
pill(s, "SPRING BOOT 3  |  ANGULAR 18  |  JAVA 17", 0.78, 4.35, 3.55, RED)
# editorial screenshot collage
add_image(s, "002.png", 7.15, 0.72, 5.55, 3.1)
add_image(s, "004.png", 7.52, 4.13, 5.2, 2.9)
box(s, 7.15, 0.72, 5.55, 3.1, None, CORAL)
box(s, 7.52, 4.13, 5.2, 2.9, None, TEAL)
text(s, "PRODUCT OVERVIEW  /  2026", 0.78, 6.77, 4.5, 0.2, 9, RGBColor(183, 180, 179), True)
footer(s, 1, True)

# 2. Positioning
s = prs.slides.add_slide(blank)
box(s, 0, 0, W, H, CREAM)
title(s, "THE JOB TO BE DONE", "Make migration work visible, ordered, and recoverable.", "The studio turns a high-risk database move into a sequence of inspectable decisions.")
box(s, 0.72, 2.25, 3.7, 3.65, BURGUNDY, BURGUNDY, True)
text(s, "FROM", 1.02, 2.63, 1.0, 0.2, 10, CORAL, True)
text(s, "Sybase ASE\nmetadata + data", 1.02, 3.02, 2.9, 0.8, 25, WHITE, True, "Georgia")
text(s, "JDBC metadata discovery\nManual-review boundaries for complex objects", 1.02, 4.55, 2.8, 0.62, 12, RGBColor(226, 220, 217))
text(s, "→", 4.57, 3.55, 0.55, 0.55, 31, RED, True, "Georgia", PP_ALIGN.CENTER)
box(s, 5.15, 2.25, 3.7, 3.65, WHITE, LINE, True)
text(s, "THROUGH", 5.47, 2.63, 1.4, 0.2, 10, RED, True)
text(s, "Migration\nStudio", 5.47, 3.02, 2.9, 0.8, 25, INK, True, "Georgia")
bullet(s, "Reusable connection profiles", 5.47, 4.37, 2.9)
bullet(s, "Ordered DDL / DML generation", 5.47, 4.82, 2.9)
bullet(s, "Checkpointed async jobs", 5.47, 5.27, 2.9)
text(s, "→", 9.0, 3.55, 0.55, 0.55, 31, RED, True, "Georgia", PP_ALIGN.CENTER)
box(s, 9.58, 2.25, 3.0, 3.65, TEAL, TEAL, True)
text(s, "TO", 9.9, 2.63, 0.8, 0.2, 10, RGBColor(208, 244, 239), True)
text(s, "Oracle\nreadiness", 9.9, 3.02, 2.2, 0.8, 25, WHITE, True, "Georgia")
text(s, "Generated scripts\nQuantitative proof\nQualitative comparison", 9.9, 4.55, 2.1, 0.9, 13, WHITE)
footer(s, 2)

# 3. Workflow
s = prs.slides.add_slide(blank)
box(s, 0, 0, W, H, NAVY)
title(s, "ONE OPERATING MODEL", "Configure once. Generate in phases. Reconcile with evidence.", "Each stage leaves an artifact or checkpoint that supports the next decision.", True)
steps = [
    ("01", "CONFIGURE", "Source + target\nprofiles", RED),
    ("02", "DISCOVER", "Schema objects\nand selections", CORAL),
    ("03", "GENERATE", "DDL, DML, keys\nand indexes", GOLD),
    ("04", "RECONCILE", "Counts, mappings\nand record proof", TEAL),
]
for i, (num, label, desc, accent) in enumerate(steps):
    x = 0.8 + i * 3.1
    box(s, x, 2.5, 2.55, 2.65, RGBColor(44, 34, 43), RGBColor(81, 64, 72), True)
    box(s, x, 2.5, 2.55, 0.08, accent, accent)
    text(s, num, x + 0.25, 2.86, 0.55, 0.28, 12, accent, True)
    text(s, label, x + 0.25, 3.4, 2.0, 0.32, 16, WHITE, True, "Aptos")
    text(s, desc, x + 0.25, 4.05, 2.0, 0.62, 13, RGBColor(221, 218, 215))
    if i < 3:
        text(s, "→", x + 2.65, 3.55, 0.38, 0.3, 20, RGBColor(157, 143, 149), True, "Georgia", PP_ALIGN.CENTER)
text(s, "DESIGN PRINCIPLE", 0.82, 5.93, 1.55, 0.2, 9, CORAL, True)
text(s, "Keep credentials ephemeral, keep progress explicit, and keep mismatches actionable.", 2.43, 5.87, 8.8, 0.3, 16, WHITE, True, "Georgia")
footer(s, 3, True)

# 4. Build migration
s = prs.slides.add_slide(blank)
box(s, 0, 0, W, H, CREAM)
title(s, "MIGRATION FACTORY", "Generate the right script at the right time.", "Base DDL, data inserts, and deferred constraints are separated to protect execution order.")
add_image(s, "002.png", 7.12, 2.02, 5.5, 3.08)
box(s, 7.12, 2.02, 5.5, 3.08, None, LINE)
text(s, "THE PHASED OUTPUT", 0.82, 2.35, 2.2, 0.2, 10, RED, True)
items = [
    ("01", "Base DDL", "Tables and views, without dependent keys or indexes."),
    ("02", "DML inserts", "Literal INSERT statements, checkpointed by table."),
    ("03", "Deferred dependencies", "Indexes, primary keys, and foreign keys after data."),
]
for i, (num, head, desc) in enumerate(items):
    y = 2.82 + i * 0.86
    text(s, num, 0.84, y, 0.42, 0.22, 10, RED, True)
    text(s, head, 1.42, y - 0.02, 2.2, 0.25, 15, INK, True)
    text(s, desc, 1.42, y + 0.28, 5.1, 0.36, 11.5, MUTED)
pill(s, "ASYNC JOBS", 0.84, 5.72, 1.2, NAVY)
pill(s, "CHECKPOINTS", 2.17, 5.72, 1.35, TEAL)
pill(s, "RESUME / RETRY", 3.65, 5.72, 1.45, BURGUNDY)
text(s, "The UI refreshes job data every five seconds and shows processed, failed, and completed units.", 7.12, 5.48, 5.45, 0.45, 13, INK, True, "Georgia")
footer(s, 4)

# 5. Connections
s = prs.slides.add_slide(blank)
box(s, 0, 0, W, H, WHITE)
title(s, "CONNECTION CONTROL", "Reusable profiles without persistent passwords.", "Connection metadata lives in file-based H2; database and keystore passwords stay in process memory.")
add_image(s, "001.png", 6.8, 2.0, 5.82, 3.2)
box(s, 6.8, 2.0, 5.82, 3.2, None, LINE)
box(s, 0.82, 2.15, 5.2, 3.2, RGBColor(248, 239, 241), RGBColor(248, 239, 241), True)
text(s, "SECURITY POSTURE", 1.15, 2.52, 2.2, 0.2, 10, RED, True)
bullet(s, "Profiles store endpoint, schema, TLS, and authentication selection.", 1.15, 3.0, 4.45, INK, RED, 13)
bullet(s, "Users do not enter JDBC URLs; the backend constructs them.", 1.15, 3.65, 4.45, INK, RED, 13)
bullet(s, "Passwords are never written to H2, logs, manifests, or scripts.", 1.15, 4.3, 4.45, INK, RED, 13)
text(s, "Operational implication: credentials must be re-entered after a JAR restart before resume, retry, or reconciliation.", 0.85, 5.77, 11.6, 0.45, 13, MUTED, False, "Aptos", PP_ALIGN.CENTER, MSO_ANCHOR.TOP, True)
footer(s, 5)

# 6. Reconciliation
s = prs.slides.add_slide(blank)
box(s, 0, 0, W, H, CREAM)
title(s, "RECONCILIATION", "Prove structure first. Then prove records.", "Quantitative reports expose inventory and count mismatches; qualitative runs compare mapped business records.")
add_image(s, "004.png", 0.82, 2.14, 5.72, 3.17)
add_image(s, "005.png", 6.82, 2.14, 5.72, 3.17)
box(s, 0.82, 2.14, 5.72, 3.17, None, LINE)
box(s, 6.82, 2.14, 5.72, 3.17, None, LINE)
pill(s, "QUANTITATIVE", 0.98, 5.54, 1.35, TEAL)
text(s, "Object inventory + COUNT(*)", 2.48, 5.59, 3.3, 0.22, 13, INK, True)
pill(s, "QUALITATIVE", 6.98, 5.54, 1.2, BURGUNDY)
text(s, "Mapped record-level comparison", 8.34, 5.59, 3.5, 0.22, 13, INK, True)
footer(s, 6)

# 7. Evidence library
s = prs.slides.add_slide(blank)
box(s, 0, 0, W, H, NAVY)
title(s, "EVIDENCE LIBRARY", "Reports remain reviewable after the run is over.", "Immutable timestamped JSON and XLSX pairs make results reloadable, shareable, and auditable.", True)
add_image(s, "003.png", 6.65, 1.95, 5.95, 3.32)
box(s, 6.65, 1.95, 5.95, 3.32, None, RGBColor(102, 92, 100))
text(s, "REPORT CONTENT", 0.84, 2.18, 2.2, 0.2, 10, CORAL, True)
for i, value in enumerate(["Summary and source / target header", "Object counts and mismatch details", "Record counts with pagination", "Stored XLSX workbook for download"]):
    bullet(s, value, 0.86, 2.72 + i * 0.61, 5.25, WHITE, CORAL, 14)
text(s, "Reload reads existing JSON without querying either database.", 0.86, 5.62, 5.1, 0.42, 14, WHITE, True, "Georgia")
footer(s, 7, True)

# 8. Architecture
s = prs.slides.add_slide(blank)
box(s, 0, 0, W, H, CREAM)
title(s, "IMPLEMENTATION SHAPE", "A focused application boundary with room to scale.", "Angular owns the operator experience; Spring Boot owns orchestration, persistence, and filesystem artifacts.")
# architecture bands
layers = [
    ("OPERATOR EXPERIENCE", "Angular 18 standalone UI  ·  persistent side menu  ·  five-second job refresh", BURGUNDY, WHITE),
    ("APPLICATION SERVICES", "Spring Boot 3  ·  REST APIs  ·  bounded worker pools  ·  memory-only secrets", RED, WHITE),
    ("PERSISTENCE + ARTIFACTS", "Flyway-managed H2  ·  JSON / XLSX reports  ·  generated DDL / DML directory tree", TEAL, WHITE),
    ("CONNECTED SYSTEMS", "SAP jConnect / Sybase ASE source  ↔  Oracle target JDBC drivers", NAVY, WHITE),
]
for i, (head, desc, fill, fg) in enumerate(layers):
    y = 2.18 + i * 0.9
    box(s, 1.0, y, 11.25, 0.68, fill, fill, True)
    text(s, head, 1.3, y + 0.11, 2.65, 0.22, 10, CORAL if fill == BURGUNDY else (RGBColor(207, 247, 241) if fill == TEAL else WHITE), True)
    text(s, desc, 4.05, y + 0.11, 7.75, 0.35, 12.5, fg, False)
    if i < len(layers) - 1:
        text(s, "↓", 6.48, y + 0.69, 0.3, 0.18, 13, MUTED, True, "Georgia", PP_ALIGN.CENTER)
footer(s, 8)

# 9. Boundaries and next steps
s = prs.slides.add_slide(blank)
box(s, 0, 0, W, H, BURGUNDY)
title(s, "DELIVERY REALITY", "Strong foundation. Explicit next frontier.", "The initial implementation is usable today, with clear manual-review boundaries for production hardening.", True)
box(s, 0.82, 2.22, 5.45, 3.35, RGBColor(72, 29, 43), RGBColor(109, 57, 70), True)
text(s, "WORKING NOW", 1.15, 2.58, 2.0, 0.2, 10, CORAL, True)
for i, v in enumerate(["Tables, columns, types, INSERTs", "Indexes, primary and foreign keys", "Quantitative and qualitative reports", "Resume, retry, pagination, XLSX export"]):
    bullet(s, v, 1.15, 3.05 + i * 0.51, 4.55, WHITE, CORAL, 12.5)
box(s, 6.7, 2.22, 5.8, 3.35, CREAM, CREAM, True)
text(s, "NEXT FRONTIER", 7.03, 2.58, 2.0, 0.2, 10, RED, True)
for i, v in enumerate(["View bodies, sequences, triggers, procedures", "Synonyms, grants, partitions, computed columns", "Bulk export and chunk-level restart", "Production-grade binary and LOB conversion"]):
    bullet(s, v, 7.03, 3.05 + i * 0.51, 4.85, INK, RED, 12.5)
text(s, "RECOMMENDED NEXT STEP", 0.84, 6.15, 2.2, 0.2, 9, CORAL, True)
text(s, "Validate against a representative ASE schema, then prioritize the manual-review object classes that matter to the first production wave.", 3.03, 6.08, 9.5, 0.38, 14, WHITE, True, "Georgia")
footer(s, 9, True)

prs.core_properties.title = "Oracle Migration Studio - Project Overview"
prs.core_properties.subject = "Sybase ASE to Oracle migration script generation and reconciliation"
prs.core_properties.author = "Migration Studio"
prs.core_properties.keywords = "Sybase, Oracle, migration, reconciliation, Spring Boot, Angular"
prs.save(OUTPUT)
print(OUTPUT)
